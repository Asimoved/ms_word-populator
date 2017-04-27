# -*- coding: utf-8 -*-

import os
import re
import time
from shutil import copyfile
from docx import Document

def copy_and_populate_template(template_path, dict_data, copy_to_folder=None):
    copied_file = copy_file_with_timestamp(template_path, copy_to_folder_path=copy_to_folder)
    populate_word_form(copied_file, dict_data)

def populate_word_form(file_path, form_data):
    # TODO: loop through tables as well, as they are not considered as paragraphs.

    def find_and_replace_placeholders(paragraph):
        replaced_text = replace_text_placeholders(paragraph.text, form_data)
        paragraph.text = replaced_text

    try:
        doc = Document(file_path)
        for paragraph in doc.paragraphs:
            find_and_replace_placeholders(paragraph)
        doc.save(file_path)
    except:
        print 'File in use'
    return file_path



def replace_text_placeholders(text, form_data, ph_prefix='{!', ph_suffix='!}'):
    # TODO: take care of text encoding to accomodate special characters
    placeholders = list_placeholders(text)
    for placeholder in placeholders:
        placeholder_data = '%s%s%s' % (ph_prefix, placeholder, ph_suffix)
        if placeholder in form_data:
            text = text.replace(placeholder_data, form_data[placeholder])

    return text

#@query
def copy_file_with_timestamp(file_to_copy_path, copy_to_folder_path=None):
    template_name = os.path.split(file_to_copy_path)[-1]
    if not copy_to_folder_path:
        copy_to_folder_path = os.path.split(file_to_copy_path)[0]
    form_path = os.path.join(copy_to_folder_path, '%s_%s' % (time.strftime("%Y%m%d-%H%M%S"), template_name))
    copyfile(file_to_copy_path, form_path)
    return form_path

#@command
def create_folder():
    order_forms_folder = os.path.join(current.request.folder, 'uploads', 'order_forms', str(order_id))
    if not os.path.exists(order_forms_folder):
        os.makedirs(order_forms_folder)

#@query
def parse_filename_from_path(file_path):
    return os.path.split(file_path)[-1]

#@query
def list_placeholders(container_string, ph_prefix='{!', ph_suffix='!}'):
    regex = '(?<=%s)(.*?)(?=%s)' % (ph_prefix, ph_suffix)
    return re.findall(regex, container_string)
